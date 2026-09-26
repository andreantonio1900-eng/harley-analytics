from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
import pandas as pd
import typer

from app.ssp_storage import DEFAULT_SSP_DB, upsert_ssp_snapshot
from app.ssp_theft import (
    MONTH_NAMES_PT,
    fetch_ssp_official_export,
    fetch_ssp_official_options,
    load_ssp_vehicle_base_from_path,
    normalize_key,
    prepare_ssp_vehicle_dataframe,
)

app = typer.Typer(help="Backfill mensal da base SSP de veiculos subtraidos para DuckDB")

KIND_OPTIONS = {
    "furto": ["Furto de veiculo"],
    "roubo": ["Roubo de veiculo"],
    "all": ["Furto de veiculo", "Roubo de veiculo"],
}


def iter_months(start: str, end: str) -> list[pd.Timestamp]:
    start_ts = pd.Timestamp(start).to_period("M").to_timestamp()
    end_ts = pd.Timestamp(end).to_period("M").to_timestamp()
    if end_ts < start_ts:
        raise typer.BadParameter("`--end` precisa ser maior ou igual a `--start`.")
    return [period.to_timestamp() for period in pd.period_range(start_ts, end_ts, freq="M")]


def infer_occurrence_kind_from_workbook(workbook_path: Path) -> str:
    name = workbook_path.name.upper()
    if "SUBTRA" in name:
        return "Veiculos Subtraidos"
    if "FURT" in name:
        return "Furto de veiculo"
    if "ROUB" in name:
        return "Roubo de veiculo"
    return "Planilha SSP"


def infer_workbook_reference_year(workbook_path: Path) -> int | None:
    match = pd.Series([workbook_path.name]).str.extract(r"(20\d{2})", expand=False).iloc[0]
    if pd.isna(match):
        return None
    return int(match)


def is_valid_workbook_timestamp(ts: pd.Timestamp, reference_year: int | None) -> bool:
    if pd.isna(ts):
        return False
    if reference_year is None:
        return 2000 <= int(ts.year) <= 2100
    return reference_year - 1 <= int(ts.year) <= reference_year


def find_vehicle_sheet_name(workbook_path: Path) -> str:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            if normalize_key(sheet_name).startswith("veiculos"):
                return sheet_name
    finally:
        workbook.close()
    raise typer.BadParameter(f"Nenhuma aba VEICULOS_* encontrada em {workbook_path}")


def stream_workbook_month_frames(workbook_path: Path) -> dict[pd.Timestamp, pd.DataFrame]:
    sheet_name = find_vehicle_sheet_name(workbook_path)
    reference_year = infer_workbook_reference_year(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        header_map = {str(value).strip(): idx for idx, value in enumerate(header) if value is not None}
        source_aliases = {
            "DATA_OCORRENCIA_BO": ["DATA_OCORRENCIA_BO"],
            "RUBRICA": ["RUBRICA"],
            "DESCR_OCORRENCIA_VEICULO": ["DESCR_OCORRENCIA_VEICULO"],
            "DESCR_TIPO_VEICULO": ["DESCR_TIPO_VEICULO"],
            "DESCR_MARCA_VEICULO": ["DESCR_MARCA_VEICULO"],
            "ANO_FABRICACAO": ["ANO_FABRICACAO"],
            "ANO_MODELO": ["ANO_MODELO"],
            "PLACA_VEICULO": ["PLACA_VEICULO"],
            "DESC_COR_VEICULO": ["DESC_COR_VEICULO", "DESCR_COR_VEICULO"],
            "CIDADE": ["CIDADE", "NOME_MUNICIPIO_CIRC", "NOME_MUNICIPIO"],
            "BAIRRO": ["BAIRRO"],
            "NOME_DELEGACIA_CIRC": ["NOME_DELEGACIA_CIRC", "NOME_DELEGACIA"],
            "MES": ["MES", "MES_REGISTRO_BO"],
            "ANO": ["ANO", "ANO_REGISTRO_BO"],
        }
        resolved_columns: dict[str, int] = {}
        missing: list[str] = []
        for canonical_name, aliases in source_aliases.items():
            match = next((alias for alias in aliases if alias in header_map), None)
            if match is None:
                missing.append(canonical_name)
                continue
            resolved_columns[canonical_name] = header_map[match]
        if missing:
            raise typer.BadParameter(
                "Colunas obrigatorias ausentes na planilha: " + ", ".join(missing)
            )

        buckets: dict[pd.Timestamp, list[dict[str, object]]] = {}
        for row in rows:
            record = {column: row[column_idx] for column, column_idx in resolved_columns.items()}
            data_fato = record.get("DATA_OCORRENCIA_BO")
            if pd.isna(data_fato) or data_fato in {None, "", "NULL"}:
                continue

            ts = pd.Timestamp(data_fato).to_period("M").to_timestamp()
            if not is_valid_workbook_timestamp(ts, reference_year):
                continue
            buckets.setdefault(ts, []).append(record)

        return {
            competencia: pd.DataFrame(records, columns=list(source_aliases.keys()))
            for competencia, records in sorted(buckets.items())
        }
    finally:
        workbook.close()


@app.command()
def run(
    kind: str = typer.Option("all", "--kind", help="furto, roubo ou all"),
    start: str = typer.Option(..., "--start", help="Mes inicial, ex.: 2026-01-01"),
    end: str = typer.Option(..., "--end", help="Mes final, ex.: 2026-07-01"),
    department: str = typer.Option("0", "--department", help="Circunscricao SSP, 0 = Todos"),
    db: str = typer.Option(str(DEFAULT_SSP_DB), "--db", help="Banco DuckDB de destino"),
    force_refresh: bool = typer.Option(False, "--force-refresh", help="Baixa novamente da SSP mesmo se houver cache"),
    workbook: list[str] = typer.Option(
        None,
        "--workbook",
        help="Caminho de uma ou mais planilhas locais anuais da SSP para ingestao mes a mes.",
    ),
):
    db_path = str(Path(db).expanduser())
    workbook_paths = [Path(value).expanduser().resolve() for value in (workbook or [])]

    if workbook_paths:
        total_loaded = 0
        for workbook_path in workbook_paths:
            if not workbook_path.exists():
                raise typer.BadParameter(f"Planilha nao encontrada: {workbook_path}")

            typer.echo(f"\nPlanilha local: {workbook_path}")
            occurrence_kind = infer_occurrence_kind_from_workbook(workbook_path)
            try:
                month_frames = stream_workbook_month_frames(workbook_path)
                typer.echo(f"Meses detectados={len(month_frames)} | modo=streaming")
            except Exception as exc:
                typer.echo(f"Streaming falhou ({exc}); usando leitura tabular padrao.")
                df, meta = load_ssp_vehicle_base_from_path(str(workbook_path))
                if df.empty:
                    typer.echo("Nenhuma linha valida encontrada na planilha.")
                    continue
                typer.echo(
                    f"Linhas carregadas={len(df):,} | sheet={meta['sheet_name']} | header_row={meta['header_row']}".replace(",", ".")
                )
                month_frames = {
                    pd.Timestamp(competencia): month_df.copy()
                    for competencia, month_df in df.groupby("competencia", dropna=True)
                    if not pd.isna(competencia)
                }

            for ts, month_source_df in month_frames.items():
                month_label = f"{MONTH_NAMES_PT[int(ts.month)]}/{int(ts.year)}"
                month_df = (
                    prepare_ssp_vehicle_dataframe(month_source_df)
                    if "competencia" not in month_source_df.columns
                    else month_source_df
                )
                loaded = upsert_ssp_snapshot(
                    df=month_df,
                    occurrence_kind=occurrence_kind,
                    query_year=int(ts.year),
                    query_month=int(ts.month),
                    department_value=str(department),
                    department_label="Todos",
                    source_label=str(workbook_path),
                    db_path=db_path,
                )
                total_loaded += loaded
                typer.echo(
                    f"[{occurrence_kind}] {month_label} | upsert ok | linhas={loaded:,}".replace(",", ".")
                )

        typer.echo(f"\nBackfill concluido | linhas inseridas/atualizadas={total_loaded:,}".replace(",", "."))
        return

    selected_kinds = KIND_OPTIONS.get(kind.lower())
    if not selected_kinds:
        raise typer.BadParameter("`--kind` deve ser um de: furto, roubo, all.")

    months = iter_months(start, end)

    department_label = department
    try:
        options = fetch_ssp_official_options(selected_kinds[0])
        department_map = {value: label for value, label in options["departments"]}
        department_label = department_map.get(str(department), str(department))
    except Exception:
        department_label = str(department)

    typer.echo(f"Destino DuckDB: {db_path}")
    typer.echo(f"Circunscricao: {department_label} ({department})")
    typer.echo(f"Competencias: {months[0].date()} -> {months[-1].date()}")

    total_loaded = 0
    for ts in months:
        for occurrence_kind in selected_kinds:
            month_label = f"{MONTH_NAMES_PT[int(ts.month)]}/{int(ts.year)}"
            typer.echo(f"\n[{occurrence_kind}] {month_label} | baixando + parseando")
            df, meta = fetch_ssp_official_export(
                occurrence_kind=occurrence_kind,
                year=int(ts.year),
                month=int(ts.month),
                department_value=str(department),
                force_refresh=force_refresh,
            )
            typer.echo(
                f"[{occurrence_kind}] {month_label} | linhas={len(df):,}".replace(",", ".")
            )
            loaded = upsert_ssp_snapshot(
                df=df,
                occurrence_kind=occurrence_kind,
                query_year=int(ts.year),
                query_month=int(ts.month),
                department_value=str(department),
                department_label=department_label,
                source_label=str(meta["source_label"]),
                db_path=db_path,
            )
            total_loaded += loaded
            typer.echo(
                f"[{occurrence_kind}] {month_label} | upsert ok | linhas={loaded:,}".replace(",", ".")
            )

    typer.echo(f"\nBackfill concluido | linhas inseridas/atualizadas={total_loaded:,}".replace(",", "."))


if __name__ == "__main__":
    app()
